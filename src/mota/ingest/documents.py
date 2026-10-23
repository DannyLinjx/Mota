"""Requirement source import and text extraction (local only)."""

from __future__ import annotations

import hashlib
import mimetypes
import shutil
from pathlib import Path
from typing import Any

from mota.errors import InvalidRequest, NotFound
from mota.projects.store import ProjectStore
from mota.workspace.atomic import read_json, write_json_atomic

MAX_BYTES = 20 * 1024 * 1024
SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".png", ".jpg", ".jpeg"}


def import_source(store: ProjectStore, project_id: str, file_path: str) -> dict[str, Any]:
    path = Path(file_path).expanduser()
    if not path.exists() or not path.is_file():
        raise InvalidRequest("Selected file does not exist.", code="SOURCE_INVALID")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise InvalidRequest(f"Unsupported file type {suffix}.", code="SOURCE_UNSUPPORTED")
    size = path.stat().st_size
    if size > MAX_BYTES:
        raise InvalidRequest("File exceeds the 20 MiB limit.", code="SOURCE_TOO_LARGE")
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    source_id = "source_" + digest[:10]
    pdir = store.project_dir(project_id)
    sources_dir = pdir / "sources"
    sources_dir.mkdir(parents=True, exist_ok=True)
    meta_path = sources_dir / f"{source_id}.json"
    if meta_path.exists():
        return read_json(meta_path)
    target = sources_dir / f"{source_id}{suffix}"
    target.write_bytes(data)
    sections: list[dict[str, Any]] = []
    warnings: list[str] = []
    kind = suffix.lstrip(".")
    pages = 0
    if suffix in (".md", ".txt"):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("utf-8", errors="replace")
            warnings.append("File was not valid UTF-8; replacement characters were used.")
        sections = _split_text(text)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(target))
            pages = len(reader.pages)
            for index, page in enumerate(reader.pages):
                content = page.extract_text() or ""
                sections.append({"section_id": f"page-{index + 1}", "page": index + 1,
                                 "text": content})
                if len(content.strip()) < 20:
                    warnings.append(f"Page {index + 1} has little or no extractable text "
                                    f"(possibly scanned).")
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"PDF extraction failed: {exc}")
            kind = "pdf_unreadable"
    elif suffix == ".docx":
        try:
            from docx import Document

            document = Document(str(target))
            blocks: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
            text = "\n\n".join(blocks)
            sections = _split_text(text)
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"DOCX extraction failed: {exc}")
            kind = "docx_unreadable"
    else:
        warnings.append("Image sources need a vision-capable model; no text was extracted here.")
        sections = []
        kind = "image"

    status = "extracted"
    if kind in ("pdf_unreadable", "docx_unreadable", "image") or not any(
            s.get("text", "").strip() for s in sections):
        status = "needs_attention"
    meta = {
        "schema_version": 1,
        "source_id": source_id,
        "project_id": project_id,
        "filename": path.name,
        "stored_name": target.name,
        "kind": kind,
        "bytes": size,
        "sha256": digest,
        "mime": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        "pages": pages,
        "section_count": len(sections),
        "sections": sections,
        "status": status,
        "warnings": warnings,
    }
    write_json_atomic(meta_path, meta)
    return _public(meta)


def list_sources(store: ProjectStore, project_id: str) -> list[dict[str, Any]]:
    sources_dir = store.project_dir(project_id) / "sources"
    if not sources_dir.exists():
        return []
    out = []
    for path in sorted(sources_dir.glob("source_*.json")):
        data = read_json(path, default=None)
        if data:
            out.append(_public(data))
    return out


def read_section(store: ProjectStore, project_id: str, source_id: str,
                 section_ids: list[str] | None = None) -> dict[str, Any]:
    path = store.project_dir(project_id) / "sources" / f"{source_id}.json"
    data = read_json(path, default=None)
    if not data:
        raise NotFound(f"Source {source_id} not found.", code="SOURCE_NOT_FOUND")
    sections = data.get("sections", [])
    if section_ids:
        sections = [s for s in sections if s.get("section_id") in section_ids]
    text = "\n\n".join(s.get("text", "") for s in sections)
    return {"source_id": source_id, "filename": data["filename"], "text": text[:60000],
            "sections": [{"section_id": s["section_id"], "page": s.get("page")}
                         for s in sections],
            "chars": len(text), "truncated": len(text) > 60000,
            "warnings": data.get("warnings", [])}


def _split_text(text: str) -> list[dict[str, Any]]:
    sections = []
    for index, block in enumerate([b.strip() for b in text.split("\n\n") if b.strip()]):
        sections.append({"section_id": f"section-{index + 1}", "text": block})
    return sections


def _public(meta: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in meta.items() if k != "sections"} | {
        "preview": (meta.get("sections", [{}])[0].get("text", "")[:400]
                    if meta.get("sections") else "")}
