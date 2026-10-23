from __future__ import annotations

import json
import time

import pytest

from mota.app.config import ConfigService
from mota.app.service import ApplicationService
from mota.errors import MotaError
from tests.conftest import login, make_user, new_project, wait_run


def test_source_import_markdown_text_pdf_and_docx(service, tmp_path):
    session = make_user(service, "docs")
    project, _ = new_project(service, session, "Docs board")

    markdown = tmp_path / "requirements.md"
    markdown.write_text("# Power\n\nInput 5 V, output 3.3 V at 50 mA.", encoding="utf-8")
    result = service.dispatch("source.import", {
        "project_id": project, "file_path": str(markdown)}, session)
    assert result["status"] == "extracted"
    assert result["section_count"] >= 2
    listed = service.dispatch("source.list", {"project_id": project}, session)
    assert listed["sources"][0]["filename"] == "requirements.md"

    text = tmp_path / "notes.txt"
    text.write_text("LED blinks about once a second.", encoding="utf-8")
    assert service.dispatch("source.import", {
        "project_id": project, "file_path": str(text)}, session)["status"] == "extracted"

    from docx import Document

    docx_path = tmp_path / "spec.docx"
    document = Document()
    document.add_paragraph("I2C temperature sensor at address 0x48.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Supply"
    table.cell(0, 1).text = "3.3 V"
    document.save(docx_path)
    docx_result = service.dispatch("source.import", {
        "project_id": project, "file_path": str(docx_path)}, session)
    assert docx_result["status"] == "extracted"

    from pypdf import PdfWriter

    pdf_path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    with pdf_path.open("wb") as handle:
        writer.write(handle)
    pdf_result = service.dispatch("source.import", {
        "project_id": project, "file_path": str(pdf_path)}, session)
    assert pdf_result["pages"] == 2
    assert pdf_result["status"] == "needs_attention"
    assert any("scanned" in warning or "extractable" in warning
               for warning in pdf_result["warnings"])

    with pytest.raises(MotaError) as exc:
        service.dispatch("source.import", {
            "project_id": project, "file_path": str(tmp_path / "missing.md")}, session)
    assert exc.value.code == "SOURCE_INVALID"


def test_same_series_package_change_requires_reverification():
    from mota.catalog.templates import get_template
    from mota.generation.board import run_build
    from mota.validation import electrical

    result = run_build(get_template("power_3v3"), {})
    circuit = result.circuit.model_copy(deep=True)
    u1 = next(c for c in circuit.components if c.reference == "U1")
    u1.footprint_id = "DIP-8_W7.62mm_Mota"
    check = electrical.check_pin_pad_mapping(circuit)
    assert check.status == "FAIL"
    assert "does not match" in check.message


def test_project_data_survives_process_restart(data_root, service):
    session = make_user(service, "persist")
    project, _ = new_project(service, session, "Persisted")
    final = service.dispatch("run.start", {
        "project_id": project,
        "expected_revision_id": service.dispatch(
            "project.open", {"project_id": project}, session)["current_revision_id"],
        "input": {"kind": "template", "template_id": "i2c_temperature",
                  "parameters": {}, "accept_template_defaults": True},
        "idempotency_key": "persist_build"}, session)
    wait_run(service, session, final["run_id"])

    restarted = ApplicationService(data_root)
    again = login(restarted, "persist")
    opened = restarted.dispatch("project.open", {"project_id": project}, again)
    assert opened["component_count"] == 5
    assert opened["current_revision_id"] == "rev_000002"
    scenes = restarted.dispatch("scene.get", {
        "project_id": project,
        "design_version": {"revision_id": opened["current_revision_id"]},
        "view": "pcb2d"}, again)
    assert scenes["objects"]


def test_catalog_sync_task_and_model_probe_task(service, monkeypatch, tmp_path):
    session = make_user(service, "tasks")
    started = service.dispatch("catalog.sync", {"source": "builtin"}, session)
    assert started["task_id"].startswith("task_")
    deadline = time.time() + 10
    data = None
    while time.time() < deadline:
        data = service.dispatch("task.inspect", {"task_id": started["task_id"]}, session)
        if data["task"]["state"] in ("SUCCEEDED", "FAILED"):
            break
        time.sleep(0.05)
    assert data["task"]["state"] == "SUCCEEDED"
    types = [event["type"] for event in data["events"]]
    assert "task.started" in types and "task.completed" in types
    assert data["task"]["result"]["snapshot_id"] == "mota_builtin_v1"

    def fake_probe(binding, api_key, suites, progress=None):
        if progress:
            progress({"stage": "probe_start", "message": "Probing"})
        return {
            "profile_id": binding.profile_id,
            "binding": binding.model_dump(),
            "suites": [{"suite": suite, "status": "PASS", "requests": 1,
                        "latency_ms": 5, "detail": "stub"} for suite in suites],
            "usage_total": {"input_tokens": 3, "output_tokens": 1, "total_tokens": 4},
            "ok": True, "vision_supported": True,
            "billing_note": "stub",
        }

    import mota.agent.probe as probe_module
    monkeypatch.setattr(probe_module, "run_probe", fake_probe)
    service.secrets.set_profile_key(session.scope(), "openai-main", "sk-test",
                                    session_only=True)
    cfg = service.dispatch("config.get", {}, session)["config"]
    probe = service.dispatch("config.probe_model", {
        "expected_config_version": cfg["config_version"],
        "profile_id": "openai-main", "test_suite": ["text"]}, session)
    deadline = time.time() + 10
    probe_data = None
    while time.time() < deadline:
        probe_data = service.dispatch("task.inspect", {"task_id": probe["task_id"]},
                                      session)
        if probe_data["task"]["state"] in ("SUCCEEDED", "FAILED"):
            break
        time.sleep(0.05)
    assert probe_data["task"]["state"] == "SUCCEEDED"
    assert probe_data["task"]["result"]["usage_total"]["input_tokens"] == 3
    assert probe_data["task"]["result"]["suites"][0]["suite"] == "text"
    cfg_after = service.dispatch("config.get", {}, session)["config"]
    assert cfg_after["model_verified"] is True
    assert cfg_after["api_key_configured"] is True
    assert cfg_after["secret_storage"] == "session_only"
    assert "sk-test" not in json.dumps(cfg_after)


def test_local_catalog_sync_with_unsupported_entries(data_root, tmp_path):
    from mota.catalog.sync import sync_local_tree

    lib = tmp_path / "libs"
    lib.mkdir()
    (lib / "Test.kicad_sym").write_text(
        '(kicad_symbol_lib (version 20231120) (generator test)'
        ' (symbol "Widget" (property "Reference" "U" (at 0 0 0))'
        '  (symbol "Widget_1_1" (pin passive line (at 0 0 0) (length 2.54)'
        '   (name "1" (effects (font (size 1 1))))'
        '   (number "1" (effects (font (size 1 1))))))))', encoding="utf-8")
    (lib / "Broken.kicad_sym").write_text("(not-a-library)", encoding="utf-8")
    coverage = sync_local_tree(lib)
    assert coverage.indexed >= 1
    assert coverage.unsupported >= 1
    assert coverage.failures
    from mota.catalog.sync import catalog_status
    status = catalog_status()
    assert any(s.get("snapshot_id") == coverage.snapshot_id
               for s in status["snapshots"])


def test_cross_user_artifact_access_is_denied(service):
    from tests.conftest import build_template

    alice = make_user(service, "artalice")
    bob = make_user(service, "artbob")
    project, _ = new_project(service, alice, "Alice art")
    data = build_template(service, alice, project, "power_3v3")
    revision = data["run"]["committed_revision_id"]
    opened = service.dispatch("project.open", {"project_id": project}, alice)
    with pytest.raises(MotaError) as exc:
        service.dispatch("artifact.read", {
            "project_id": project,
            "design_version": {"revision_id": revision},
            "artifact_id": opened["project"]["thumbnail_artifact_id"],
            "offset": 0, "length": 64}, bob)
    assert exc.value.code in ("PROJECT_NOT_FOUND", "ARTIFACT_ACCESS_DENIED")
